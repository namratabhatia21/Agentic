"""Parse uploaded files into a list of normalized records (lower_snake_case keys)."""

import csv
import io
import re


def _norm(key: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", key.strip().lower()).strip("_")


def parse(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.DictReader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "") for h in next(it)]
        rows = [dict(zip(header, r, strict=False)) for r in it]
    elif name.endswith(".json"):
        import json

        loaded = json.loads(data)
        rows = loaded if isinstance(loaded, list) else [loaded]
    else:
        raise ValueError(f"Unsupported file type: {filename} (use .csv, .xlsx or .json)")
    out = []
    for row in rows:
        rec = {_norm(str(k)): ("" if v is None else str(v).strip()) for k, v in row.items() if k}
        if any(rec.values()):
            out.append(rec)
    return out
